from flask import Flask, render_template, request, jsonify, session
from werkzeug.security import generate_password_hash, check_password_hash
import os
import psycopg2
import psycopg2.extras
from psycopg2 import pool  # <-- IMPORTACIÓN DEL POOL AGREGADA
import datetime
import json
from decimal import Decimal, ROUND_HALF_UP

app = Flask(__name__)
# Llave secreta necesaria para manejar las sesiones de los usuarios
app.secret_key = os.environ.get('SECRET_KEY', 'super-secret-key-dashboard-2024')

# --- CONFIGURACIÓN DE SESIONES PARA LA NUBE (RENDER) ---
app.config['SESSION_COOKIE_SECURE'] = True      # Obligatorio para HTTPS en Render
app.config['SESSION_COOKIE_HTTPONLY'] = True
app.config['SESSION_COOKIE_SAMESITE'] = 'Lax'       # La app es same-origin; Lax evita pérdidas de sesión por políticas cross-site innecesarias
app.config['SESSION_COOKIE_PATH'] = '/'
app.config['PERMANENT_SESSION_LIFETIME'] = datetime.timedelta(days=7)

# --- INICIALIZAR EL POOL GLOBAL DE CONEXIONES ---
db_pool = None
database_url = os.environ.get('DATABASE_URL')

if database_url:
    try:
        # Mantiene entre 1 y 20 conexiones vivas listas para usar (ideal para Flask)
        db_pool = psycopg2.pool.ThreadedConnectionPool(1, 20, database_url)
    except Exception as e:
        print(f"Error al iniciar el pool de conexiones: {e}")

# --- WRAPPER PARA COMPATIBILIDAD SQLITE -> POSTGRESQL ---
class PostgresConnWrapper:
    def __init__(self, conn):
        self.conn = conn
        self.conn.autocommit = False # Mantiene el control manual de transacciones

    def execute(self, query, args=None):
        cur = self.conn.cursor(cursor_factory=psycopg2.extras.DictCursor)
        # Reemplazar placeholders de SQLite (?) a PostgreSQL (%s)
        query = query.replace('?', '%s')
        if args:
            cur.execute(query, args)
        else:
            cur.execute(query)
        return cur

    def commit(self):
        self.conn.commit()

    def close(self):
        # AHORA DEVOLVEMOS LA CONEXIÓN AL POOL EN LUGAR DE CERRARLA FÍSICAMENTE
        if db_pool:
            self.conn.rollback() # Limpia cualquier transacción a medias por seguridad
            db_pool.putconn(self.conn)
        else:
            self.conn.close() # Fallback por si el pool no cargó

def get_db_connection():
    if db_pool:
        # Tomamos una conexión instantánea del pool
        conn = db_pool.getconn()
        return PostgresConnWrapper(conn)
    else:
        # Fallback clásico si algo falla con el pool
        database_url = os.environ.get('DATABASE_URL')
        if not database_url:
            raise ValueError("No se encontró la variable de entorno DATABASE_URL")
        conn = psycopg2.connect(database_url)
        return PostgresConnWrapper(conn)

def safe_alter(conn, query):
    """Ejecuta un ALTER TABLE y si falla hace rollback seguro"""
    try:
        conn.execute(query)
        conn.commit()
    except psycopg2.Error:
        conn.conn.rollback()

def init_db():
    conn = get_db_connection()
    try:
        conn.execute('''CREATE TABLE IF NOT EXISTS clientes (id SERIAL PRIMARY KEY, documento TEXT, nombre TEXT, telefono TEXT, correo TEXT, pais TEXT, estado TEXT, municipio TEXT, direccion_entrega TEXT, punto_referencia TEXT, coordenadas TEXT, tipo_envio TEXT, fecha_registro TEXT, registrado_por TEXT)''')
        conn.execute('''CREATE TABLE IF NOT EXISTS proveedores (id SERIAL PRIMARY KEY, nombre TEXT, correo TEXT, telefono TEXT, direccion TEXT, tipo TEXT, fecha_registro TEXT, registrado_por TEXT)''')
        conn.execute('''CREATE TABLE IF NOT EXISTS almacenes (id SERIAL PRIMARY KEY, nombre TEXT, ubicacion TEXT, fecha_registro TEXT, registrado_por TEXT)''')
        conn.execute('''CREATE TABLE IF NOT EXISTS categorias (id SERIAL PRIMARY KEY, nombre TEXT, descripcion TEXT, fecha_registro TEXT, registrado_por TEXT)''')
        conn.execute('''CREATE TABLE IF NOT EXISTS productos (id SERIAL PRIMARY KEY, categoria_id INTEGER, proveedor_id INTEGER, descripcion TEXT, unidad_medida TEXT, stock_minimo INTEGER, precio_usd REAL, estado TEXT, codigo_barras TEXT, foto TEXT, fecha_registro TEXT, registrado_por TEXT)''')
        conn.execute('''CREATE TABLE IF NOT EXISTS movimientos (id SERIAL PRIMARY KEY, consecutivo TEXT, fecha_registro TEXT, tipo TEXT, producto_id INTEGER, cantidad REAL, costo_unitario REAL, almacen_origen_id INTEGER, almacen_destino_id INTEGER, motivo TEXT, documento TEXT, registrado_por TEXT)''')
        conn.execute('''CREATE TABLE IF NOT EXISTS ventas (id SERIAL PRIMARY KEY, consecutivo TEXT, fecha_registro TEXT, cliente_nombre TEXT, cliente_telefono TEXT, direccion_entrega TEXT, total_eur REAL, total_bs REAL DEFAULT 0, tasa_bcv_euro_aplicada REAL, tasa_binance_aplicada REAL, porcentaje_brecha_aplicado REAL, estado TEXT, registrado_por TEXT, metodo_pago TEXT)''')
        conn.execute('''CREATE TABLE IF NOT EXISTS detalle_nota_entrega (id SERIAL PRIMARY KEY, consecutivo TEXT, producto_id INTEGER, cantidad REAL, descuento REAL, precio_unitario_euro_snapshot REAL, subtotal_euro_snapshot REAL, total_euro_snapshot REAL, precio_unitario_bs_snapshot REAL, subtotal_bs_snapshot REAL, total_bs_snapshot REAL)''')
        
        # Módulo Devoluciones y Notas de Crédito
        conn.execute('''CREATE TABLE IF NOT EXISTS notas_credito (id SERIAL PRIMARY KEY, consecutivo TEXT, fecha_registro TEXT, consecutivo_origen TEXT, cliente_nombre TEXT, total_eur REAL, total_bs REAL, motivo TEXT, registrado_por TEXT, saldo_usado_eur REAL DEFAULT 0, saldo_usado_bs REAL DEFAULT 0, estado TEXT DEFAULT 'DISPONIBLE')''')
        conn.execute('''CREATE TABLE IF NOT EXISTS detalle_nota_credito (id SERIAL PRIMARY KEY, consecutivo_nc TEXT, producto_id INTEGER, cantidad REAL, precio_eur REAL, precio_bs REAL, subtotal_eur REAL, subtotal_bs REAL)''')
        
        conn.execute('''CREATE TABLE IF NOT EXISTS historico_precios_dia (fecha TEXT PRIMARY KEY, json_data TEXT)''')
        conn.execute('''CREATE TABLE IF NOT EXISTS historico_tasas (id SERIAL PRIMARY KEY, fecha TEXT, hora TEXT, dolar_bcv REAL, binance REAL, bybit REAL, dolar_promedio REAL, euro_bcv REAL, zelle REAL, paypal REAL, brecha REAL, registrado_por TEXT)''')
        conn.execute('''CREATE TABLE IF NOT EXISTS historico_coberturas (id SERIAL PRIMARY KEY, fecha_registro TEXT, rango_evaluado TEXT, fecha_pico_maximo TEXT, porcentaje_cobertura REAL, factor_proteccion REAL, registrado_por TEXT, estado TEXT)''')
        conn.execute('''CREATE TABLE IF NOT EXISTS configuracion (clave TEXT PRIMARY KEY, valor TEXT)''')
        
        # --- TABLA DE USUARIOS ---
        conn.execute('''CREATE TABLE IF NOT EXISTS usuarios (id SERIAL PRIMARY KEY, nombre TEXT, usuario TEXT UNIQUE, contrasena TEXT, activo BOOLEAN DEFAULT TRUE, es_admin BOOLEAN DEFAULT FALSE, permisos TEXT, protegido BOOLEAN DEFAULT FALSE, fecha_registro TEXT)''')
        
        conn.commit()

        # Migraciones automáticas (Usando safe_alter)
        safe_alter(conn, "ALTER TABLE productos ADD COLUMN precio_usd REAL DEFAULT 0")
        safe_alter(conn, "ALTER TABLE ventas ADD COLUMN metodo_pago TEXT DEFAULT ''")
        safe_alter(conn, "ALTER TABLE ventas ADD COLUMN total_bs REAL DEFAULT 0")
        safe_alter(conn, "ALTER TABLE notas_credito ADD COLUMN saldo_usado_eur REAL DEFAULT 0")
        safe_alter(conn, "ALTER TABLE notas_credito ADD COLUMN estado TEXT DEFAULT 'DISPONIBLE'")

        # Mantiene la columna histórica alineada con la fórmula oficial de brecha.
        conn.execute('''UPDATE historico_tasas
                        SET brecha = CASE
                            WHEN COALESCE(euro_bcv, 0) > 0 THEN (binance / euro_bcv) - 1
                            ELSE 0
                        END''')
        
        columnas_clientes = ['correo', 'pais', 'estado', 'municipio', 'punto_referencia', 'coordenadas', 'tipo_envio']
        for col in columnas_clientes:
            safe_alter(conn, f"ALTER TABLE clientes ADD COLUMN {col} TEXT DEFAULT ''")

        # Configuración por defecto
        conn.execute("INSERT INTO configuracion (clave, valor) VALUES ('permitir_descuentos', 'true') ON CONFLICT (clave) DO NOTHING")
        
        # Almacenes Automáticos
        conn.execute("INSERT INTO almacenes (id, nombre, ubicacion, fecha_registro, registrado_por) VALUES (9999, 'Devoluciones por Venta (BLOQUEADO)', 'Sistema Automático', CURRENT_TIMESTAMP, 'Sistema') ON CONFLICT (id) DO NOTHING")
        conn.execute("INSERT INTO almacenes (id, nombre, ubicacion, fecha_registro, registrado_por) VALUES (9998, 'Almacén de Merma (BLOQUEADO)', 'Sistema Automático', CURRENT_TIMESTAMP, 'Sistema') ON CONFLICT (id) DO NOTHING")
        
        if conn.execute("SELECT COUNT(*) FROM historico_coberturas").fetchone()[0] == 0:
            conn.execute('''INSERT INTO historico_coberturas (fecha_registro, rango_evaluado, fecha_pico_maximo, porcentaje_cobertura, factor_proteccion, registrado_por, estado) VALUES (CURRENT_TIMESTAMP, 'Inicial', 'N/A', 0.20, 1.20, 'Sistema', 'ACTIVO')''')
        
        # Insertar Usuario Fantasma (Admin) si no existe
        admin = conn.execute("SELECT * FROM usuarios WHERE usuario = 'admin'").fetchone()
        if not admin:
            hashed = generate_password_hash('admin')
            conn.execute("INSERT INTO usuarios (nombre, usuario, contrasena, activo, es_admin, permisos, protegido, fecha_registro) VALUES (?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)",
                         ('Administrador Principal', 'admin', hashed, True, True, '[]', True))

        conn.commit()
    finally:
        conn.close()

if os.environ.get('DATABASE_URL'):
    init_db()

def safe_float(val):
    try: return float(val) if val else 0.0
    except: return 0.0

def safe_int(val):
    try: return int(val) if val else 0
    except: return 0

def tiene_permiso_en_sesion(permiso):
    """Comprueba permisos de la sesión para operaciones sensibles."""
    if not session.get('usuario_id'):
        return False
    if session.get('es_admin'):
        return True
    return permiso in session.get('permisos', [])

def puede_gestionar_parametros():
    return tiene_permiso_en_sesion('parametros')

def puede_agregar_tasa():
    return puede_gestionar_parametros() or tiene_permiso_en_sesion('agregar_tasa')

def es_administrador_actual():
    return bool(session.get('usuario_id') and session.get('es_admin'))

def precio_eur_comercial(precio_usd, factor_proteccion):
    valor = (Decimal(str(precio_usd or 0)) * Decimal(str(factor_proteccion or 1)))
    valor = valor.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    entero_siguiente = valor.to_integral_value(rounding=ROUND_HALF_UP)
    if valor < entero_siguiente and (entero_siguiente - valor) <= Decimal('0.01'):
        valor = entero_siguiente
    return float(valor)

def respuesta_sin_permiso():
    return jsonify({'error': 'No es posible realizar esta operación.'}), 403

def es_ultimo_administrador_activo(conn, usuario_id):
    """Indica si el usuario es el único administrador activo restante."""
    usuario = conn.execute('SELECT es_admin, activo FROM usuarios WHERE id = ?', (usuario_id,)).fetchone()
    if not usuario or not usuario['es_admin'] or not usuario['activo']:
        return False
    cantidad = conn.execute('SELECT COUNT(*) FROM usuarios WHERE es_admin = TRUE AND activo = TRUE').fetchone()[0]
    return cantidad <= 1

# ----------------- RUTAS DE SESIÓN Y AUTENTICACIÓN -----------------

@app.route('/api/auth/login', methods=['POST'])
def api_login():
    try:
        # force=True soluciona el error si el frontend no envía las cabeceras correctamente
        data = request.get_json(force=True) 
        usuario = data.get('usuario')
        contrasena = data.get('contrasena')
        
        conn = get_db_connection()
        try:
            row = conn.execute("SELECT * FROM usuarios WHERE usuario = ?", (usuario,)).fetchone()
            if row and check_password_hash(row['contrasena'], contrasena):
                if not row['activo']:
                    return jsonify({'error': 'El usuario se encuentra inactivo.'}), 401
                    
                # Parsear permisos desde el JSON almacenado
                try: permisos_list = json.loads(row['permisos']) if row['permisos'] else []
                except: permisos_list = []
                
                session.permanent = True
                session['usuario_id'] = row['id']
                session['nombre'] = row['nombre']
                session['usuario'] = row['usuario']
                session['es_admin'] = row['es_admin']
                session['permisos'] = permisos_list
                
                return jsonify({
                    'id': row['id'], 'nombre': row['nombre'], 'usuario': row['usuario'],
                    'es_admin': row['es_admin'], 'permisos': permisos_list
                })
            return jsonify({'error': 'Usuario o contraseña inválidos.'}), 401
        finally:
            conn.close()
    except Exception as e:
        # Si ocurre CUALQUIER error interno, ahora lo mostrará en tu pantalla
        return jsonify({'error': f"Error interno detectado: {str(e)}"}), 500

@app.route('/api/auth/sesion', methods=['GET'])
def api_sesion():
    if 'usuario_id' in session:
        return jsonify({
            'id': session['usuario_id'],
            'nombre': session['nombre'],
            'usuario': session['usuario'],
            'es_admin': session['es_admin'],
            'permisos': session.get('permisos', [])
        })
    return jsonify({'autenticado': False}), 401

@app.route('/api/auth/logout', methods=['POST'])
def api_logout():
    session.clear()
    return jsonify({'status': 'ok'})

# ----------------- RUTAS PRINCIPALES DEL SISTEMA -----------------

@app.route('/ping')
def ping_keep_alive():
    return "OK", 200

@app.route('/')
def index(): 
    return render_template('index.html')

@app.route('/api/resumen', methods=['GET'])
def api_resumen():
    conn = get_db_connection()
    try:
        c = conn.execute('SELECT COUNT(*) FROM clientes').fetchone()[0]
        p = conn.execute('SELECT COUNT(*) FROM proveedores').fetchone()[0]
        prod = conn.execute('SELECT COUNT(*) FROM productos').fetchone()[0]
        v = conn.execute('SELECT COUNT(*) FROM ventas').fetchone()[0]
        pend_clientes = conn.execute("SELECT COUNT(*) FROM clientes WHERE documento = 'PENDIENTE' OR documento = ''").fetchone()[0]
        pend_stock = conn.execute("SELECT COUNT(*) FROM productos p WHERE (COALESCE((SELECT SUM(cantidad) FROM movimientos WHERE producto_id = p.id AND tipo IN ('Inventario Inicial', 'Compra', 'Devolución por venta', 'Ajuste administrativo - Entrada')), 0) - COALESCE((SELECT SUM(cantidad) FROM movimientos WHERE producto_id = p.id AND tipo IN ('Venta', 'Descarga por daño/motivo', 'Devolución por compra', 'Ajuste administrativo - Salida')), 0)) <= p.stock_minimo").fetchone()[0]
        return jsonify({'conteo': {'clientes': c, 'proveedores': p, 'productos': prod, 'ventas': v}, 'notificaciones': {'clientes_pendientes': pend_clientes, 'stock_bajo': pend_stock}})
    except Exception as e: return jsonify({'error': str(e)}), 500
    finally: conn.close()

@app.route('/api/configuracion', methods=['GET', 'POST'])
def api_configuracion():
    conn = get_db_connection()
    try:
        if request.method == 'POST':
            if not tiene_permiso_en_sesion('configuracion'):
                return respuesta_sin_permiso()
            d = request.json or {}
            if set(d.keys()) - {'permitir_descuentos'}:
                return jsonify({'error': 'La identidad y apariencia de ZUARA APP son fijas y no son editables.'}), 400
            if 'permitir_descuentos' in d:
                valor = 'true' if str(d['permitir_descuentos']).lower() == 'true' else 'false'
                conn.execute("UPDATE configuracion SET valor = ? WHERE clave = 'permitir_descuentos'", (valor,))
            conn.commit()
            fila = conn.execute("SELECT valor FROM configuracion WHERE clave = 'permitir_descuentos'").fetchone()
            return jsonify({'status': 'ok', 'configuracion': {'permitir_descuentos': fila['valor'] if fila else 'true'}})
        fila = conn.execute("SELECT valor FROM configuracion WHERE clave = 'permitir_descuentos'").fetchone()
        return jsonify({'permitir_descuentos': fila['valor'] if fila else 'true'})
    except Exception as e: return jsonify({'error': str(e)}), 500
    finally: conn.close()

@app.route('/api/stock_almacenes/<int:producto_id>', methods=['GET'])
def api_stock_almacenes(producto_id):
    conn = get_db_connection()
    try:
        query = "SELECT a.id, a.nombre, COALESCE(SUM(CASE WHEN m.almacen_destino_id = a.id AND m.tipo IN ('Inventario Inicial', 'Compra', 'Traspaso', 'Devolución por venta', 'Ajuste administrativo - Entrada') THEN m.cantidad ELSE 0 END), 0) - COALESCE(SUM(CASE WHEN m.almacen_origen_id = a.id AND m.tipo IN ('Traspaso', 'Descarga por daño/motivo', 'Devolución por compra', 'Ajuste administrativo - Salida') THEN m.cantidad ELSE 0 END), 0) as stock FROM almacenes a LEFT JOIN movimientos m ON (a.id = m.almacen_destino_id OR m.almacen_origen_id = a.id) AND m.producto_id = ? GROUP BY a.id"
        data = conn.execute(query, (producto_id,)).fetchall()
        return jsonify([dict(ix) for ix in data])
    finally: conn.close()

@app.route('/api/movimientos', methods=['POST'])
def registrar_movimiento():
    if not tiene_permiso_en_sesion('movimientos'):
        return respuesta_sin_permiso()
    usuario_actual = session.get('nombre', 'Sistema')
    conn = get_db_connection()
    try:
        d = request.json
        ahora = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        fecha_mov = d.get('fecha_registro') if d.get('fecha_registro') else ahora
        ult_mov = conn.execute('SELECT id FROM movimientos ORDER BY id DESC LIMIT 1').fetchone()
        consecutivo = f"MOV-{str((ult_mov[0] + 1) if ult_mov else 1).zfill(5)}"
        
        conn.execute('''INSERT INTO movimientos (consecutivo, fecha_registro, tipo, producto_id, cantidad, costo_unitario, almacen_origen_id, almacen_destino_id, motivo, documento, registrado_por) VALUES (?,?,?,?,?,?,?,?,?,?,?)''',
                      (consecutivo, fecha_mov, d['tipo'], d['producto_id'], safe_float(d.get('cantidad')), safe_float(d.get('costo_unitario')), safe_int(d.get('almacen_origen_id')) or None, safe_int(d.get('almacen_destino_id')) or None, d.get('motivo',''), d.get('documento',''), usuario_actual))
        if 'precio_usd' in d and d['precio_usd'] != "":
            if not es_administrador_actual():
                conn.conn.rollback()
                return respuesta_sin_permiso()
            conn.execute("UPDATE productos SET precio_usd = ? WHERE id = ?", (safe_float(d['precio_usd']), d['producto_id']))
            
        conn.commit()
        return jsonify({'status': 'ok'})
    except Exception as e: return jsonify({'error': str(e)}), 500
    finally: conn.close()

@app.route('/api/existencias/<int:producto_id>/corregir', methods=['POST'])
def corregir_ultima_carga(producto_id):
    if not es_administrador_actual():
        return respuesta_sin_permiso()
    conn = get_db_connection()
    try:
        d = request.json or {}
        movimiento_id = safe_int(d.get('movimiento_id'))
        precio_usd = d.get('precio_usd')
        precio_actual = safe_float(precio_usd) if precio_usd is not None and precio_usd != '' else None
        if precio_actual is not None and precio_actual < 0:
            return jsonify({'error': 'El precio objetivo no puede ser negativo.'}), 400

        carga = None
        if movimiento_id:
            carga = conn.execute("""SELECT id, consecutivo, fecha_registro, tipo, producto_id, cantidad, costo_unitario,
                                       almacen_destino_id, motivo, documento
                                FROM movimientos
                                WHERE id = ? AND producto_id = ? AND tipo IN ('Inventario Inicial', 'Compra')
                                FOR UPDATE""", (movimiento_id, producto_id)).fetchone()
            if not carga:
                return jsonify({'error': 'La carga indicada no existe o no es una carga corregible.'}), 404
            nueva_cantidad = safe_float(d.get('cantidad'))
            nuevo_costo = safe_float(d.get('costo_unitario'))
            if nueva_cantidad < 0 or nuevo_costo < 0:
                return jsonify({'error': 'La cantidad y el costo no pueden ser negativos.'}), 400
            if not carga['almacen_destino_id']:
                return jsonify({'error': 'La carga no tiene un almacén destino válido.'}), 400

            cantidad_anterior = float(carga['cantidad'] or 0)
            costo_anterior = float(carga['costo_unitario'] or 0)
            delta = round(nueva_cantidad - cantidad_anterior, 10)
            if delta < 0:
                stock_actual = conn.execute("""SELECT COALESCE(SUM(CASE
                    WHEN tipo IN ('Inventario Inicial', 'Compra', 'Devolución por venta', 'Ajuste administrativo - Entrada') THEN cantidad
                    WHEN tipo IN ('Venta', 'Descarga por daño/motivo', 'Devolución por compra', 'Ajuste administrativo - Salida') THEN -cantidad
                    ELSE 0 END), 0) AS stock
                    FROM movimientos WHERE producto_id = ?""", (producto_id,)).fetchone()['stock'] or 0
                if abs(delta) > float(stock_actual) + 1e-9:
                    return jsonify({'error': f'La corrección dejaría el stock físico en negativo. Stock actual: {float(stock_actual):g}; reducción solicitada: {abs(delta):g}.'}), 400
            if abs(delta) > 0.0000000001 or abs(nuevo_costo - costo_anterior) > 0.0000001:
                tipo_ajuste = 'Ajuste administrativo - Entrada' if delta >= 0 else 'Ajuste administrativo - Salida'
                cantidad_ajuste = abs(delta)
                ult_mov = conn.execute('SELECT id FROM movimientos ORDER BY id DESC LIMIT 1').fetchone()
                consecutivo = f"MOV-{str((ult_mov[0] + 1) if ult_mov else 1).zfill(5)}"
                detalle = (f"Corrección administrativa de carga {carga['consecutivo']}: cantidad {cantidad_anterior:g} -> {nueva_cantidad:g}; "
                           f"costo {costo_anterior:.2f} -> {nuevo_costo:.2f}")
                conn.execute("""INSERT INTO movimientos
                    (consecutivo, fecha_registro, tipo, producto_id, cantidad, costo_unitario,
                     almacen_origen_id, almacen_destino_id, motivo, documento, registrado_por)
                    VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
                    (consecutivo, datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"), tipo_ajuste,
                     producto_id, cantidad_ajuste, nuevo_costo,
                     carga['almacen_destino_id'] if delta < 0 else None,
                     carga['almacen_destino_id'] if delta >= 0 else None,
                     detalle, carga['documento'] or carga['consecutivo'], session.get('nombre', 'Sistema')))
        else:
            if precio_actual is None:
                return jsonify({'error': 'Este producto no tiene una carga registrada. Solo se puede editar el precio objetivo hasta registrar una carga.'}), 400

        if precio_actual is not None:
            conn.execute('UPDATE productos SET precio_usd = ? WHERE id = ?', (precio_actual, producto_id))
        conn.commit()
        return jsonify({'status': 'ok', 'ajuste': consecutivo if carga and 'consecutivo' in locals() else None,
                        'delta_cantidad': delta if carga and 'delta' in locals() else 0})
    except Exception as e:
        conn.conn.rollback()
        return jsonify({'error': str(e)}), 500
    finally:
        conn.close()

@app.route('/api/existencias', methods=['GET'])
def api_existencias():
    conn = get_db_connection()
    try:
        query = '''
            SELECT p.id, p.codigo_barras, p.descripcion, p.stock_minimo, p.unidad_medida, p.estado, 
            COALESCE(SUM(CASE WHEN m.tipo IN ('Inventario Inicial', 'Compra', 'Devolución por venta', 'Ajuste administrativo - Entrada') THEN m.cantidad ELSE 0 END), 0) - COALESCE(SUM(CASE WHEN m.tipo IN ('Venta', 'Descarga por daño/motivo', 'Devolución por compra', 'Ajuste administrativo - Salida') THEN m.cantidad ELSE 0 END), 0) as stock_fisico_total,
            COALESCE(SUM(CASE WHEN m.almacen_destino_id = 9999 THEN m.cantidad ELSE 0 END), 0) - COALESCE(SUM(CASE WHEN m.almacen_origen_id = 9999 THEN m.cantidad ELSE 0 END), 0) as stock_devoluciones,
            COALESCE(SUM(CASE WHEN m.almacen_destino_id = 9998 THEN m.cantidad ELSE 0 END), 0) - COALESCE(SUM(CASE WHEN m.almacen_origen_id = 9998 THEN m.cantidad ELSE 0 END), 0) as stock_merma,
            (SELECT costo_unitario FROM movimientos WHERE producto_id = p.id AND tipo IN ('Inventario Inicial', 'Compra', 'Ajuste administrativo - Entrada', 'Ajuste administrativo - Salida') ORDER BY id DESC LIMIT 1) as costo_unit,
            p.precio_usd,
            (SELECT id FROM movimientos WHERE producto_id = p.id AND tipo IN ('Inventario Inicial', 'Compra') ORDER BY id DESC LIMIT 1) as ultima_carga_id,
            (SELECT cantidad FROM movimientos WHERE producto_id = p.id AND tipo IN ('Inventario Inicial', 'Compra') ORDER BY id DESC LIMIT 1) as ultima_carga_cantidad,
            (SELECT fecha_registro FROM movimientos WHERE producto_id = p.id AND tipo IN ('Inventario Inicial', 'Compra') ORDER BY id DESC LIMIT 1) as ultima_carga_fecha,
            (SELECT documento FROM movimientos WHERE producto_id = p.id AND tipo IN ('Inventario Inicial', 'Compra') ORDER BY id DESC LIMIT 1) as ultima_carga_documento
            FROM productos p LEFT JOIN movimientos m ON p.id = m.producto_id GROUP BY p.id ORDER BY p.descripcion ASC
        '''
        data = conn.execute(query).fetchall()
        resultados = []
        for ix in data:
            row = dict(ix)
            row['stock_bloqueado'] = float(row['stock_devoluciones'] or 0) + float(row['stock_merma'] or 0)
            row['stock_disponible_venta'] = float(row['stock_fisico_total'] or 0) - row['stock_bloqueado']
            row['costo_unit'] = row['costo_unit'] or 0
            row['total_costo'] = float(row['stock_fisico_total'] or 0) * float(row['costo_unit'])
            resultados.append(row)
        return jsonify(resultados)
    finally: conn.close()

@app.route('/api/kardex', methods=['GET'])
def api_kardex():
    conn = get_db_connection()
    try:
        data = conn.execute('''SELECT m.*, COALESCE(p.descripcion, 'Producto Eliminado') as producto_nombre, 
                            ao.nombre as almacen_origen_nombre, ad.nombre as almacen_destino_nombre
                            FROM movimientos m LEFT JOIN productos p ON m.producto_id = p.id 
                            LEFT JOIN almacenes ao ON m.almacen_origen_id = ao.id
                            LEFT JOIN almacenes ad ON m.almacen_destino_id = ad.id
                            ORDER BY m.id DESC''').fetchall()
        return jsonify([dict(ix) for ix in data])
    finally: conn.close()

@app.route('/api/ventas', methods=['GET', 'POST'])
def api_ventas():
    usuario_actual = session.get('nombre', 'Sistema')
    conn = get_db_connection()
    try:
        if request.method == 'POST':
            ahora = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            d = request.json
            c_nombre = d.get('cliente_nombre', '').strip()
            
            if not c_nombre: return jsonify({'error': 'El nombre del cliente es obligatorio'}), 400
            if not d.get('detalles'): return jsonify({'error': 'El carrito está vacío'}), 400

            consec_venta = str(d.get('consecutivo') or '').strip()
            if not consec_venta:
                return jsonify({'error': 'No se pudo generar el consecutivo de la nota de entrega.'}), 400
            if conn.execute('SELECT id FROM ventas WHERE consecutivo = ?', (consec_venta,)).fetchone():
                return jsonify({'error': 'Ya existe una nota de entrega con ese consecutivo.'}), 409

            cantidades_por_producto = {}
            for item in d['detalles']:
                producto_id = safe_int(item.get('producto_id'))
                cantidad = safe_float(item.get('cantidad'))
                if producto_id <= 0 or cantidad <= 0:
                    return jsonify({'error': 'Cada detalle de venta debe tener un producto y una cantidad mayor a cero.'}), 400
                cantidades_por_producto[producto_id] = cantidades_por_producto.get(producto_id, 0) + cantidad

            for producto_id, cantidad_solicitada in cantidades_por_producto.items():
                stock = conn.execute('''SELECT
                                        COALESCE(SUM(CASE
                                            WHEN tipo IN ('Inventario Inicial', 'Compra', 'Devolución por venta', 'Ajuste administrativo - Entrada') THEN cantidad
                                            WHEN tipo IN ('Venta', 'Descarga por daño/motivo', 'Devolución por compra', 'Ajuste administrativo - Salida') THEN -cantidad
                                            ELSE 0
                                        END), 0)
                                        - (
                                            COALESCE(SUM(CASE WHEN almacen_destino_id IN (9998, 9999) THEN cantidad ELSE 0 END), 0)
                                            - COALESCE(SUM(CASE WHEN almacen_origen_id IN (9998, 9999) THEN cantidad ELSE 0 END), 0)
                                        ) AS disponible
                                    FROM movimientos WHERE producto_id = ?''', (producto_id,)).fetchone()
                if cantidad_solicitada > float(stock['disponible'] or 0):
                    return jsonify({'error': 'No hay existencias suficientes para completar la venta.'}), 400
            
            existe = conn.execute('SELECT id FROM clientes WHERE nombre = ?', (c_nombre,)).fetchone()
            if existe:
                conn.execute('UPDATE clientes SET documento=?, telefono=?, correo=?, pais=?, estado=?, municipio=?, direccion_entrega=?, punto_referencia=?, coordenadas=?, tipo_envio=? WHERE id=?',
                              (d.get('cliente_doc', 'PENDIENTE'), d.get('cliente_telefono', ''), d.get('cliente_correo', ''), d.get('env_pais', 'Venezuela'), d.get('env_estado', ''), d.get('env_ciudad', ''), d.get('env_direccion', ''), d.get('env_referencia', ''), d.get('env_coordenadas', ''), d.get('env_tipo', ''), existe['id']))
            else:
                conn.execute('INSERT INTO clientes (documento, nombre, telefono, correo, pais, estado, municipio, direccion_entrega, punto_referencia, coordenadas, tipo_envio, fecha_registro, registrado_por) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)',
                              (d.get('cliente_doc', 'PENDIENTE'), c_nombre, d.get('cliente_telefono', ''), d.get('cliente_correo', ''), d.get('env_pais', 'Venezuela'), d.get('env_estado', ''), d.get('env_ciudad', ''), d.get('env_direccion', ''), d.get('env_referencia', ''), d.get('env_coordenadas', ''), d.get('env_tipo', ''), ahora, usuario_actual))
            
            total_eur = safe_float(d.get('total_eur'))
            total_bs = safe_float(d.get('total_bs'))
            
            nc_id = safe_int(d.get('nc_id'))
            if nc_id:
                nota_credito = conn.execute('''SELECT id, total_eur, total_bs, saldo_usado_eur, saldo_usado_bs, estado
                                                FROM notas_credito WHERE id = ? FOR UPDATE''', (nc_id,)).fetchone()
                if not nota_credito or nota_credito['estado'] != 'DISPONIBLE':
                    return jsonify({'error': 'La nota de crédito seleccionada ya no está disponible.'}), 400
                saldo_disponible = float(nota_credito['total_eur'] or 0) - float(nota_credito['saldo_usado_eur'] or 0)
                if total_eur > saldo_disponible + 0.0001:
                    return jsonify({'error': 'El total de la venta supera el saldo disponible de la nota de crédito.'}), 400
                saldo_usado_eur = float(nota_credito['saldo_usado_eur'] or 0) + total_eur
                saldo_usado_bs = float(nota_credito['saldo_usado_bs'] or 0) + total_bs
                estado_nc = 'APLICADA' if saldo_usado_eur >= float(nota_credito['total_eur'] or 0) - 0.0001 else 'DISPONIBLE'
                conn.execute('''UPDATE notas_credito SET saldo_usado_eur = ?, saldo_usado_bs = ?, estado = ? WHERE id = ?''',
                             (saldo_usado_eur, saldo_usado_bs, estado_nc, nc_id))
            conn.execute('INSERT INTO ventas (consecutivo, fecha_registro, cliente_nombre, cliente_telefono, direccion_entrega, total_eur, total_bs, tasa_bcv_euro_aplicada, tasa_binance_aplicada, porcentaje_brecha_aplicado, estado, registrado_por, metodo_pago) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)',
                          (consec_venta, ahora, c_nombre, d.get('cliente_telefono',''), d.get('env_direccion',''), total_eur, total_bs, safe_float(d.get('tasa_bcv_euro')), safe_float(d.get('tasa_binance')), safe_float(d.get('brecha_dia')), d.get('estado_semaforo','EMITIDA'), usuario_actual, d.get('metodo_pago', '')))
            ult_mov = conn.execute('SELECT id FROM movimientos ORDER BY id DESC LIMIT 1').fetchone()
            num_m = (ult_mov[0] + 1) if ult_mov else 1
            for item in d['detalles']:
                consec_mov = f"MOV-{str(num_m).zfill(5)}"
                conn.execute('''INSERT INTO movimientos (consecutivo, fecha_registro, tipo, producto_id, cantidad, costo_unitario, documento, registrado_por) VALUES (?,?,?,?,?,?,?,?)''',
                              (consec_mov, ahora, 'Venta', item['producto_id'], safe_float(item['cantidad']), safe_float(item.get('precio_eur')), consec_venta, usuario_actual))
                num_m += 1
                
                conn.execute('''INSERT INTO detalle_nota_entrega (consecutivo, producto_id, cantidad, descuento, precio_unitario_euro_snapshot, subtotal_euro_snapshot, total_euro_snapshot, precio_unitario_bs_snapshot, subtotal_bs_snapshot, total_bs_snapshot) VALUES (?,?,?,?,?,?,?,?,?,?)''',
                              (consec_venta, item['producto_id'], item['cantidad'], item['descuento'], item['precio_eur'], item['sub_eur'], item['total_eur'], item['pre_bs'], item['sub_bs'], item['tot_bs']))
            conn.commit()
            return jsonify({'status': 'ok', 'consecutivo': consec_venta})
        
        data = conn.execute('SELECT * FROM ventas ORDER BY id DESC').fetchall()
        return jsonify([dict(ix) for ix in data])
    except Exception as e: return jsonify({'error': str(e)}), 500
    finally: conn.close()

@app.route('/api/ventas/detalles/<consecutivo>', methods=['GET'])
def get_detalles_venta(consecutivo):
    conn = get_db_connection()
    try:
        query = '''SELECT d.*, COALESCE(p.descripcion, 'Producto Eliminado') as producto_nombre, p.codigo_barras as codigo
                    FROM detalle_nota_entrega d LEFT JOIN productos p ON d.producto_id = p.id WHERE d.consecutivo = ?'''
        data = conn.execute(query, (consecutivo,)).fetchall()
        return jsonify([dict(ix) for ix in data])
    except Exception as e: return jsonify({'error': str(e)}), 500
    finally: conn.close()

@app.route('/api/clientes/notas_credito/<cliente_nombre>', methods=['GET'])
def get_notas_credito_cliente(cliente_nombre):
    conn = get_db_connection()
    try:
        query = "SELECT * FROM notas_credito WHERE cliente_nombre = ? AND estado = 'DISPONIBLE'"
        data = conn.execute(query, (cliente_nombre,)).fetchall()
        return jsonify([dict(ix) for ix in data])
    finally: conn.close()

@app.route('/api/devoluciones', methods=['POST'])
def registrar_devolucion():
    usuario_actual = session.get('nombre', 'Sistema')
    conn = get_db_connection()
    try:
        d = request.json
        ahora = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        consec_origen = d['consecutivo_origen']

        detalles_validados = []
        cantidades_solicitadas = {}
        for item in d.get('detalles', []):
            producto_id = safe_int(item.get('producto_id'))
            cantidad = safe_float(item.get('cantidad_devolver'))
            if producto_id <= 0 or cantidad <= 0:
                continue
            cantidades_solicitadas[producto_id] = cantidades_solicitadas.get(producto_id, 0) + cantidad

        if not cantidades_solicitadas:
            return jsonify({'error': 'Debe seleccionar al menos un producto con cantidad válida para devolver.'}), 400

        for producto_id, cantidad_solicitada in cantidades_solicitadas.items():
            vendido = conn.execute('''SELECT COALESCE(SUM(cantidad), 0) AS cantidad,
                                             COALESCE(MAX(precio_unitario_euro_snapshot), 0) AS precio_eur,
                                             COALESCE(MAX(precio_unitario_bs_snapshot), 0) AS precio_bs
                                      FROM detalle_nota_entrega
                                      WHERE consecutivo = ? AND producto_id = ?''', (consec_origen, producto_id)).fetchone()
            devuelto = conn.execute('''SELECT COALESCE(SUM(d.cantidad), 0) AS cantidad
                                       FROM detalle_nota_credito d
                                       INNER JOIN notas_credito n ON n.consecutivo = d.consecutivo_nc
                                       WHERE n.consecutivo_origen = ? AND d.producto_id = ?''', (consec_origen, producto_id)).fetchone()
            disponible = float(vendido['cantidad'] or 0) - float(devuelto['cantidad'] or 0)
            if cantidad_solicitada > disponible + 0.0001:
                return jsonify({'error': 'La cantidad solicitada supera la cantidad disponible para devolución.'}), 400
            detalles_validados.append({
                'producto_id': producto_id,
                'cantidad_devolver': cantidad_solicitada,
                'precio_eur': float(vendido['precio_eur'] or 0),
                'precio_bs': float(vendido['precio_bs'] or 0)
            })

        total_eur = sum(item['cantidad_devolver'] * item['precio_eur'] for item in detalles_validados)
        total_bs = sum(item['cantidad_devolver'] * item['precio_bs'] for item in detalles_validados)
        
        # Generar Consecutivo Nota de Crédito
        consec_nc = consec_origen.replace('NE', 'NC', 1)
        existe_nc = conn.execute('SELECT id FROM notas_credito WHERE consecutivo = ?', (consec_nc,)).fetchone()
        if existe_nc:
            consec_nc = f"{consec_nc}-{int(datetime.datetime.now().timestamp())}"
        
        conn.execute('''INSERT INTO notas_credito (consecutivo, fecha_registro, consecutivo_origen, cliente_nombre, total_eur, total_bs, motivo, registrado_por, estado) VALUES (?,?,?,?,?,?,?,?,?)''',
                      (consec_nc, ahora, consec_origen, d['cliente_nombre'], total_eur, total_bs, d['motivo'], usuario_actual, 'DISPONIBLE'))
        ult_mov = conn.execute('SELECT id FROM movimientos ORDER BY id DESC LIMIT 1').fetchone()
        num_m = (ult_mov[0] + 1) if ult_mov else 1
        for item in detalles_validados:
            if item['cantidad_devolver'] > 0:
                consec_mov = f"MOV-{str(num_m).zfill(5)}"
                motivo_kardex = f"Afecta a Nota: {consec_origen} | Motivo: {d['motivo']}"
                
                conn.execute('''INSERT INTO movimientos (consecutivo, fecha_registro, tipo, producto_id, cantidad, costo_unitario, almacen_destino_id, documento, registrado_por, motivo) VALUES (?,?,?,?,?,?,?,?,?,?)''',
                              (consec_mov, ahora, 'Devolución por venta', item['producto_id'], safe_float(item['cantidad_devolver']), safe_float(item['precio_eur']), 9999, consec_nc, usuario_actual, motivo_kardex))
                num_m += 1
                
                conn.execute('''INSERT INTO detalle_nota_credito (consecutivo_nc, producto_id, cantidad, precio_eur, precio_bs, subtotal_eur, subtotal_bs) VALUES (?,?,?,?,?,?,?)''',
                              (consec_nc, item['producto_id'], safe_float(item['cantidad_devolver']), safe_float(item['precio_eur']), safe_float(item.get('precio_bs', 0)), safe_float(item['cantidad_devolver']) * safe_float(item['precio_eur']), safe_float(item['cantidad_devolver']) * safe_float(item.get('precio_bs', 0))))
        conn.execute("UPDATE ventas SET estado = 'DEVUELTO PARCIAL/TOTAL' WHERE consecutivo = ?", (consec_origen,))
        conn.commit()
        return jsonify({'status': 'ok', 'consecutivo': consec_nc})
    except Exception as e: return jsonify({'error': str(e)}), 500
    finally: conn.close()

@app.route('/api/notas_credito/detalles/<consecutivo_nc>', methods=['GET'])
def get_detalles_nota_credito(consecutivo_nc):
    conn = get_db_connection()
    try:
        query = '''SELECT d.*, COALESCE(p.descripcion, 'Producto Eliminado') as producto_nombre, p.codigo_barras as codigo
                    FROM detalle_nota_credito d LEFT JOIN productos p ON d.producto_id = p.id WHERE d.consecutivo_nc = ?'''
        data = conn.execute(query, (consecutivo_nc,)).fetchall()
        return jsonify([dict(ix) for ix in data])
    except Exception as e: return jsonify({'error': str(e)}), 500
    finally: conn.close()

@app.route('/api/lista_precios_data', methods=['GET'])
def api_lista_precios_data():
    conn = get_db_connection()
    try:
        hoy = datetime.datetime.now().strftime("%Y-%m-%d")
        cob = conn.execute("SELECT porcentaje_cobertura, factor_proteccion FROM historico_coberturas WHERE estado='ACTIVO' ORDER BY id DESC LIMIT 1").fetchone()
        factor = float(cob['factor_proteccion']) if cob else 1.0
        cobertura_activa = float(cob['porcentaje_cobertura']) if cob else 0.0
        tasa = conn.execute("SELECT fecha, hora, binance, euro_bcv FROM historico_tasas WHERE fecha = ? ORDER BY hora DESC LIMIT 1", (hoy,)).fetchone()
        if not tasa:
            t_bin = 0.0; t_eur = 0.0; brecha_dia = 0.0
            estado = "FALTAN_TASAS"; etiqueta = "REGISTRE TASA DEL DÍA"
        else:
            t_bin = float(tasa['binance'] or 0)
            t_eur = float(tasa['euro_bcv'] or 0)
            brecha_dia = (t_bin / t_eur) - 1 if t_eur > 0 else 0
            if brecha_dia > cobertura_activa: estado = "MERCADO_VOLATIL"; etiqueta = "MERCADO VOLÁTIL / AJUSTAR PRECIO"
            elif brecha_dia > 0: estado = "PRECIO_SEGURO_PRECAUCION"; etiqueta = "PRECIO SEGURO"
            else: estado = "PRECIO_SEGURO_EXCELENTE"; etiqueta = "PRECIO SEGURO"
        prods = conn.execute("SELECT p.id, p.codigo_barras as codigo, c.nombre as categoria, p.descripcion, p.unidad_medida, p.precio_usd FROM productos p LEFT JOIN categorias c ON p.categoria_id = c.id WHERE p.estado='ACTIVO'").fetchall()
        resultados = []
        for p in prods:
            p_usd = float(p['precio_usd'] or 0)
            p_eur = precio_eur_comercial(p_usd, factor)
            p_bs = round(p_eur * t_eur, 2)
            resultados.append({
                'id': p['id'], 'codigo': p['codigo'], 'categoria': p['categoria'] or 'N/A', 'descripcion': p['descripcion'],
                'unidad_medida': p['unidad_medida'], 'precio_usd': p_usd, 'precio_eur': p_eur, 'precio_bs': p_bs,
                'estado_semaforo': estado, 'etiqueta_semaforo': etiqueta
            })
        data_final = {
            'tasas': {'fecha': tasa['fecha'] if tasa else hoy, 'hora': tasa['hora'] if tasa else '--:--', 'binance': t_bin, 'euro_bcv': t_eur, 'brecha': brecha_dia, 'cobertura_activa': cobertura_activa, 'registrada_hoy': bool(tasa)},
            'productos': resultados
        }
        conn.execute("INSERT INTO historico_precios_dia (fecha, json_data) VALUES (?, ?) ON CONFLICT (fecha) DO UPDATE SET json_data = EXCLUDED.json_data", (hoy, json.dumps(data_final)))
        conn.commit()
        return jsonify(data_final)
    finally: conn.close()

@app.route('/api/historico_precios/<fecha>', methods=['GET'])
def get_historico_precios(fecha):
    conn = get_db_connection()
    try:
        data = conn.execute("SELECT json_data FROM historico_precios_dia WHERE fecha = ?", (fecha,)).fetchone()
        if data: return jsonify(json.loads(data['json_data']))
        return jsonify({'error': 'No hay registros de precios para esta fecha.'}), 404
    finally: conn.close()

@app.route('/api/tasas/upload', methods=['POST'])
def upload_tasas():
    if not puede_gestionar_parametros():
        return respuesta_sin_permiso()
    if 'file' not in request.files: return jsonify({'error': 'No file'}), 400
    f = request.files['file']
    usuario_actual = session.get('nombre', 'Sistema')
    try:
        import openpyxl
        wb = openpyxl.load_workbook(f, data_only=True)
        sheet = wb.active
        conn = get_db_connection()
        inserted = 0
        for row in sheet.iter_rows(min_row=8, values_only=True):
            if row[1] is None: continue
            fecha_val = row[1]
            fecha = fecha_val.strftime('%Y-%m-%d') if hasattr(fecha_val, 'strftime') else str(fecha_val).split(' ')[0]
            hora_val = row[9]
            hora = hora_val.strftime('%H:%M:%S') if hasattr(hora_val, 'strftime') else str(hora_val)[:8]
            if len(hora) == 5: hora += ":00"
            bcv, binance, byb, prom, euro, zel, pay = [safe_float(row[i]) for i in (2, 3, 4, 5, 6, 7, 8)]
            brecha = (binance / euro) - 1 if euro > 0 else 0
            conn.execute('''INSERT INTO historico_tasas (fecha, hora, dolar_bcv, binance, bybit, dolar_promedio, euro_bcv, zelle, paypal, brecha, registrado_por) VALUES (?,?,?,?,?,?,?,?,?,?,?)''',
               (fecha, hora, bcv, binance, byb, prom, euro, zel, pay, brecha, usuario_actual))
            inserted += 1
        conn.commit()
        return jsonify({'status': 'ok', 'inserted': inserted})
    except Exception as e: return jsonify({'error': str(e)}), 500
    finally:
        try: conn.close()
        except: pass

@app.route('/api/tasas/brecha-maxima', methods=['GET'])
def obtener_brecha_maxima():
    """Devuelve la fila real que contiene la mayor brecha del rango indicado."""
    if not puede_gestionar_parametros():
        return respuesta_sin_permiso()

    fecha_inicio = request.args.get('fecha_inicio', '').strip()
    fecha_fin = request.args.get('fecha_fin', '').strip()
    if not fecha_inicio or not fecha_fin:
        return jsonify({'error': 'Las fechas de inicio y fin son obligatorias.'}), 400
    if fecha_inicio > fecha_fin:
        return jsonify({'error': 'La fecha de inicio no puede ser mayor a la fecha final.'}), 400

    conn = get_db_connection()
    try:
        fila = conn.execute('''
            SELECT id, fecha, hora, binance, euro_bcv,
                   ((binance::double precision / euro_bcv::double precision) - 1) AS brecha
            FROM historico_tasas
            WHERE fecha >= ?
              AND fecha <= ?
              AND COALESCE(euro_bcv, 0) > 0
            ORDER BY ((binance::double precision / euro_bcv::double precision) - 1) DESC,
                     fecha ASC, hora ASC, id ASC
            LIMIT 1
        ''', (fecha_inicio, fecha_fin)).fetchone()

        if not fila:
            return jsonify({'error': 'No hay tasas válidas con EURO BCV mayor a cero en el rango seleccionado.'}), 404

        return jsonify(dict(fila))
    except Exception as e:
        return jsonify({'error': str(e)}), 500
    finally:
        conn.close()

# ----------------- RUTAS DINÁMICAS (CRUD GENÉRICO) -----------------
@app.route('/api/<tabla>', methods=['GET', 'POST'])
def api_tabla(tabla):
    if tabla not in ['clientes', 'proveedores', 'almacenes', 'categorias', 'productos', 'tasas', 'coberturas', 'usuarios', 'notas_credito']:
        return jsonify({'error': 'Tabla no permitida'}), 403
    return api_crud(tabla, request)

@app.route('/api/<tabla>/<int:id>', methods=['PUT', 'DELETE'])
def api_tabla_id(tabla, id):
    if tabla not in ['clientes', 'proveedores', 'almacenes', 'categorias', 'productos', 'tasas', 'coberturas', 'ventas', 'usuarios']:
        return jsonify({'error': 'Tabla no permitida'}), 403
    return api_crud(tabla, request, id)

def api_crud(tabla, request, id=None):
    conn = get_db_connection()
    usuario_actual = session.get('nombre', 'Sistema')
    try:
        ahora = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        tabla_db = 'historico_tasas' if tabla == 'tasas' else tabla
        tabla_db = 'historico_coberturas' if tabla == 'coberturas' else tabla_db

        if tabla == 'usuarios' and not tiene_permiso_en_sesion('usuarios'):
            return respuesta_sin_permiso()

        # Un usuario con el permiso especial puede registrar una tasa desde el
        # dashboard, pero no consultar, editar ni borrar información sensible.
        if tabla == 'tasas':
            if request.method == 'POST' and not puede_agregar_tasa():
                return respuesta_sin_permiso()
            if request.method != 'POST' and not puede_gestionar_parametros():
                return respuesta_sin_permiso()

        # El historial de coberturas es parte del módulo crítico de parámetros.
        if tabla == 'coberturas' and not puede_gestionar_parametros():
            return respuesta_sin_permiso()
        
        if request.method == 'GET':
            query = f'SELECT * FROM {tabla_db} ORDER BY id DESC'
            if tabla_db == 'productos': 
                query = '''SELECT p.*, c.nombre as categoria_nombre, prov.nombre as proveedor_nombre FROM productos p LEFT JOIN categorias c ON p.categoria_id = c.id LEFT JOIN proveedores prov ON p.proveedor_id = prov.id ORDER BY p.id DESC'''
            elif tabla_db == 'historico_tasas': 
                query = 'SELECT * FROM historico_tasas ORDER BY fecha DESC, hora DESC'
            elif tabla_db == 'usuarios':
                query = 'SELECT id, nombre, usuario, activo, es_admin, permisos, protegido, fecha_registro FROM usuarios ORDER BY id ASC'
                
            data = conn.execute(query).fetchall()
            return jsonify([dict(ix) for ix in data])
            
        elif request.method == 'POST':
            d = request.json
            if tabla_db == 'clientes': 
                conn.execute('INSERT INTO clientes (documento, nombre, telefono, correo, pais, estado, municipio, direccion_entrega, punto_referencia, coordenadas, tipo_envio, fecha_registro, registrado_por) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)', 
                             (d.get('documento','PENDIENTE'), d.get('nombre',''), d.get('telefono',''), d.get('correo',''), d.get('pais',''), d.get('estado',''), d.get('municipio',''), d.get('direccion_entrega',''), d.get('punto_referencia',''), d.get('coordenadas',''), d.get('tipo_envio',''), ahora, usuario_actual))
            elif tabla_db == 'proveedores': 
                conn.execute('INSERT INTO proveedores (nombre, correo, telefono, direccion, tipo, fecha_registro, registrado_por) VALUES (?,?,?,?,?,?,?)', 
                             (d.get('nombre',''), d.get('correo',''), d.get('telefono',''), d.get('direccion',''), d.get('tipo',''), ahora, usuario_actual))
            elif tabla_db == 'almacenes': 
                conn.execute('INSERT INTO almacenes (nombre, ubicacion, fecha_registro, registrado_por) VALUES (?,?,?,?)', 
                             (d.get('nombre',''), d.get('ubicacion',''), ahora, usuario_actual))
            elif tabla_db == 'categorias': 
                conn.execute('INSERT INTO categorias (nombre, descripcion, fecha_registro, registrado_por) VALUES (?,?,?,?)', 
                             (d.get('nombre',''), d.get('descripcion',''), ahora, usuario_actual))
            elif tabla_db == 'productos': 
                conn.execute('INSERT INTO productos (categoria_id, proveedor_id, descripcion, unidad_medida, stock_minimo, precio_usd, estado, codigo_barras, foto, fecha_registro, registrado_por) VALUES (?,?,?,?,?,?,?,?,?,?,?)', 
                             (safe_int(d.get('categoria_id')) or None, safe_int(d.get('proveedor_id')) or None, d.get('descripcion',''), d.get('unidad_medida',''), safe_int(d.get('stock_minimo')), safe_float(d.get('precio_usd', 0)), d.get('estado','ACTIVO'), d.get('codigo_barras',''), d.get('foto',''), ahora, usuario_actual))
            elif tabla_db == 'historico_tasas':
                binance = safe_float(d.get('binance'))
                euro = safe_float(d.get('euro_bcv'))
                brecha = (binance / euro) - 1 if euro > 0 else 0
                conn.execute('''INSERT INTO historico_tasas (fecha, hora, dolar_bcv, binance, bybit, dolar_promedio, euro_bcv, zelle, paypal, brecha, registrado_por) VALUES (?,?,?,?,?,?,?,?,?,?,?)''',
                              (d['fecha'], d['hora'], safe_float(d.get('dolar_bcv')), binance, safe_float(d.get('bybit')), safe_float(d.get('dolar_promedio')), euro, safe_float(d.get('zelle')), safe_float(d.get('paypal')), brecha, usuario_actual))
            elif tabla_db == 'historico_coberturas':
                conn.execute('''INSERT INTO historico_coberturas (fecha_registro, rango_evaluado, fecha_pico_maximo, porcentaje_cobertura, factor_proteccion, registrado_por, estado) VALUES (?,?,?,?,?,?,?)''',
                              (ahora, d.get('rango_evaluado',''), d.get('fecha_pico_maximo',''), safe_float(d.get('porcentaje_cobertura')), safe_float(d.get('factor_proteccion')), usuario_actual, d.get('estado', 'ACTIVO')))
            elif tabla_db == 'usuarios':
                hashed = generate_password_hash(d['contrasena'])
                conn.execute("INSERT INTO usuarios (nombre, usuario, contrasena, activo, es_admin, permisos, fecha_registro) VALUES (?, ?, ?, ?, ?, ?, ?)",
                             (d['nombre'], d['usuario'], hashed, d.get('activo', True), d.get('es_admin', False), json.dumps(d.get('permisos', [])), ahora))
            conn.commit()
            return jsonify({'status': 'ok'})
            
        elif request.method == 'PUT':
            d = request.json
            if tabla_db == 'productos' and 'precio_usd' in d:
                actual = conn.execute('SELECT precio_usd FROM productos WHERE id = ?', (id,)).fetchone()
                if actual and abs(float(actual['precio_usd'] or 0) - safe_float(d.get('precio_usd'))) > 0.0000001 and not es_administrador_actual():
                    return respuesta_sin_permiso()
            if tabla_db == 'clientes': 
                conn.execute('UPDATE clientes SET documento=?, nombre=?, telefono=?, correo=?, pais=?, estado=?, municipio=?, direccion_entrega=?, punto_referencia=?, coordenadas=?, tipo_envio=? WHERE id=?', 
                             (d.get('documento','PENDIENTE'), d.get('nombre',''), d.get('telefono',''), d.get('correo',''), d.get('pais',''), d.get('estado',''), d.get('municipio',''), d.get('direccion_entrega',''), d.get('punto_referencia',''), d.get('coordenadas',''), d.get('tipo_envio',''), id))
            elif tabla_db == 'proveedores': 
                conn.execute('UPDATE proveedores SET nombre=?, correo=?, telefono=?, direccion=?, tipo=? WHERE id=?', 
                             (d.get('nombre',''), d.get('correo',''), d.get('telefono',''), d.get('direccion',''), d.get('tipo',''), id))
            elif tabla_db == 'almacenes': 
                conn.execute('UPDATE almacenes SET nombre=?, ubicacion=? WHERE id=?', 
                             (d.get('nombre',''), d.get('ubicacion',''), id))
            elif tabla_db == 'categorias': 
                conn.execute('UPDATE categorias SET nombre=?, descripcion=? WHERE id=?', 
                             (d.get('nombre',''), d.get('descripcion',''), id))
            elif tabla_db == 'productos': 
                conn.execute('UPDATE productos SET categoria_id=?, proveedor_id=?, descripcion=?, unidad_medida=?, stock_minimo=?, precio_usd=?, estado=?, codigo_barras=?, foto=? WHERE id=?', 
                             (safe_int(d.get('categoria_id')) or None, safe_int(d.get('proveedor_id')) or None, d.get('descripcion',''), d.get('unidad_medida',''), safe_int(d.get('stock_minimo')), safe_float(d.get('precio_usd', 0)), d.get('estado','ACTIVO'), d.get('codigo_barras',''), d.get('foto',''), id))
            elif tabla_db == 'historico_tasas':
                conn.execute('UPDATE historico_tasas SET fecha=?, hora=?, dolar_bcv=?, binance=?, bybit=?, dolar_promedio=?, euro_bcv=?, zelle=?, paypal=?, brecha=? WHERE id=?', 
                             (d['fecha'], d['hora'], safe_float(d.get('dolar_bcv')), safe_float(d.get('binance')), safe_float(d.get('bybit')), safe_float(d.get('dolar_promedio')), safe_float(d.get('euro_bcv')), safe_float(d.get('zelle')), safe_float(d.get('paypal')), (safe_float(d.get('binance')) / safe_float(d.get('euro_bcv'))) - 1 if safe_float(d.get('euro_bcv')) > 0 else 0, id))
            elif tabla_db == 'historico_coberturas':
                conn.execute('''UPDATE historico_coberturas
                                SET rango_evaluado=?, fecha_pico_maximo=?, porcentaje_cobertura=?,
                                    factor_proteccion=?, estado=?
                                WHERE id=?''',
                             (d.get('rango_evaluado', ''), d.get('fecha_pico_maximo', ''),
                              safe_float(d.get('porcentaje_cobertura')),
                              safe_float(d.get('factor_proteccion')), d.get('estado', 'ACTIVO'), id))
            elif tabla_db == 'usuarios':
                if es_ultimo_administrador_activo(conn, id) and not (d.get('es_admin', False) and d.get('activo', True)):
                    return jsonify({'error': 'Debe permanecer al menos un administrador activo en el sistema.'}), 400
                if d.get('contrasena'):
                    hashed = generate_password_hash(d['contrasena'])
                    conn.execute("UPDATE usuarios SET nombre=?, usuario=?, contrasena=?, activo=?, es_admin=?, permisos=? WHERE id=?",
                                 (d['nombre'], d['usuario'], hashed, d.get('activo', True), d.get('es_admin', False), json.dumps(d.get('permisos', [])), id))
                else:
                    conn.execute("UPDATE usuarios SET nombre=?, usuario=?, activo=?, es_admin=?, permisos=? WHERE id=?",
                                 (d['nombre'], d['usuario'], d.get('activo', True), d.get('es_admin', False), json.dumps(d.get('permisos', [])), id))
            conn.commit()
            return jsonify({'status': 'ok'})
            
        elif request.method == 'DELETE':
            if tabla_db == 'ventas':
                venta = conn.execute("SELECT consecutivo FROM ventas WHERE id=?", (id,)).fetchone()
                if not venta:
                    return jsonify({'error': 'La venta indicada no existe.'}), 404
                conn.execute("DELETE FROM movimientos WHERE documento=?", (venta['consecutivo'],))
                conn.execute("DELETE FROM detalle_nota_entrega WHERE consecutivo=?", (venta['consecutivo'],))
                conn.execute('DELETE FROM ventas WHERE id=?', (id,))
            elif tabla_db == 'usuarios':
                if es_ultimo_administrador_activo(conn, id):
                    return jsonify({'error': 'No se puede eliminar el último administrador activo.'}), 400
                conn.execute('DELETE FROM usuarios WHERE id=? AND protegido=FALSE', (id,))
            else:
                conn.execute(f'DELETE FROM {tabla_db} WHERE id=?', (id,))
            conn.commit()
            return jsonify({'status': 'ok'})
    except Exception as e: return jsonify({'error': str(e)}), 500
    finally: conn.close()

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port)
